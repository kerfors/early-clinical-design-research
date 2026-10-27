"""Build tpp_template_comparison.ipynb with the json module.

Usage: python3 build_tpp_notebook.py <output.ipynb> <source_dir_for_notebook>
"""
import json
import sys


def markdown_cell(text):
    lines = text.strip("\n").split("\n")
    return {"cell_type": "markdown", "metadata": {}, "source": [line + "\n" for line in lines]}


def code_cell(text):
    lines = text.strip("\n").split("\n")
    return {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [],
            "source": [line + "\n" for line in lines]}


OUTPUT_PATH = sys.argv[1]
SOURCE_DIR = sys.argv[2]
cells = []

cells.append(markdown_cell("""
# Target Product Profile templates — structural comparison

Extracts the structure of four public TPP templates/examples into one long table and compares their column models and attributes.

| Key | Source | Date | Nature |
|---|---|---|---|
| `FDA2007` | FDA draft guidance *Target Product Profile — A Strategic Development Process Tool*, Appendix C template | March 2007 (draft; no longer in FDA's guidance list) | blank template, label-anchored |
| `WOAH2026` | WOAH *Target Product Profile – Template for animal vaccines* | Feb 2026 | blank template, attribute table |
| `PAVEGT` | NCATS PaVe-GT *Target Product Profile (TPP) Template* (AAV gene therapy) | 2025 (upload path 2025/07) | blank template, attribute table |
| `SEED_*` | NIH SEED *Example TPP* for small molecule, biologic, vaccine, cell & gene therapy | Dec 2023 | **filled examples** (illustrative values) |

All content comes from the files. The SEED documents are filled examples, so their cell values are kept in `example_value_*` columns and are **not** template guidance. NIH SEED states that the examples are illustrative.
"""))

cells.append(code_cell(f"""
import re
import urllib.request
from pathlib import Path

import docx
import pandas as pd
import pdfplumber
from docx.oxml.ns import qn
from docx.table import _Cell

SOURCE_DIR = Path("{SOURCE_DIR}")
OUTPUT_DIR = Path("output_tpp")
OUTPUT_DIR.mkdir(exist_ok=True)
pd.set_option("display.max_colwidth", 80)
pd.set_option("display.width", 220)

SOURCES = {{
    "FDA2007": ("fda_tpp_2007.pdf",
                "https://fda.report/media/72566/Target-Product-Profile----A-Strategic-Development-Process-Tool.pdf"),
    "WOAH2026": ("woah_tpp_template.pdf",
                 "https://www.woah.org/app/uploads/2026/02/a-202601-ahg-tpp-template.pdf"),
    "PAVEGT": ("pavegt_tpp_template.docx",
               "https://pave-gt.ncats.nih.gov/wp-content/uploads/2025/07/PaVe-GT-Target-Product-Profile-Template-508.docx"),
    "SEED_SMALL_MOLECULE": ("seed_Example-TPP-Small-Molecule-Drug.pdf",
                            "https://seed.nih.gov/sites/default/files/2023-12/Example-TPP-Small-Molecule-Drug.pdf"),
    "SEED_BIOLOGIC": ("seed_Example-TPP-Biologic.pdf",
                      "https://seed.nih.gov/sites/default/files/2023-12/Example-TPP-Biologic.pdf"),
    "SEED_VACCINE": ("seed_Example-TPP-Vaccine.pdf",
                     "https://seed.nih.gov/sites/default/files/2023-12/Example-TPP-Vaccine.pdf"),
    "SEED_CGT": ("seed_Example-TPP-for-Cell-and-Gene-Therapy.pdf",
                 "https://seed.nih.gov/sites/default/files/2023-12/Example-TPP-for-Cell-and-Gene-Therapy.pdf"),
}}
"""))

cells.append(markdown_cell("""
## Get the source files

Downloads any file missing from `SOURCE_DIR`. The FDA guidance is no longer on fda.gov (its page returns "Page Not Found"), so the copy at fda.report is used. Its Federal Register notice is 72 FR 15141, 30 March 2007, docket 2007D-0089.
"""))

cells.append(code_cell("""
SOURCE_DIR.mkdir(exist_ok=True)
for key, (name, url) in SOURCES.items():
    target = SOURCE_DIR / name
    if not target.exists():
        print(f"downloading {key}: {url}")
        urllib.request.urlretrieve(url, target)
    print(f"{key:22s} {target.stat().st_size:>9,d} bytes  {name}")
"""))

cells.append(markdown_cell("""
## FDA 2007 — Appendix C

Appendix C (PDF pages 16–25) is two-column text, not a table. The parser works on word positions:

- Line numbers sit at x < 60 and are dropped. Page header and footer lines are dropped.
- A **section** starts on a line whose first word (at x ≈ 72) is a section number such as `1`, `8`, `8.1` or `12.3`.
- Within a main section, the `Target` / `Annotations` header line sets the column split at the x-position of `Annotations`. Words to the left are **target prompt** text and words to the right are **annotation prompt** text, up to `Comments:`.
- Subsections (`8.1 Pregnancy …`) are full-width text and are kept as `guidance_text`. Text after `:` in a subsection heading, or after ` — ` in a section heading (e.g. *15 References — Can include …*), goes into `guidance_text`.
- The milestone box is read from the first Appendix C page with `extract_tables()`. Its header spans three rows, which are combined column-wise.
"""))

cells.append(code_cell("""
FDA_SECTION_RE = re.compile(r"^\\d{1,2}(?:\\.\\d)?$")


def fda_lines(pdf_path, first_page, last_page):
    \"\"\"Return list of lines: dict(page, top, words[list of word dicts]).\"\"\"
    lines = []
    with pdfplumber.open(pdf_path) as pdf:
        for pno in range(first_page, last_page + 1):
            page = pdf.pages[pno - 1]
            words = [w for w in page.extract_words(keep_blank_chars=False)
                     if w["top"] > 65 and w["bottom"] < page.height - 40]
            words.sort(key=lambda w: (round(w["top"]), w["x0"]))
            current, current_top = [], None
            for w in words:
                if current_top is None or abs(w["top"] - current_top) > 3:
                    if current:
                        lines.append({"page": pno, "top": current_top, "words": current})
                    current, current_top = [w], w["top"]
                else:
                    current.append(w)
            if current:
                lines.append({"page": pno, "top": current_top, "words": current})
    for line in lines:
        line["words"] = [w for w in line["words"] if not (w["x0"] < 60 and w["text"].isdigit())]
    return [line for line in lines if line["words"]]


def text_of(words):
    return " ".join(w["text"] for w in words).strip()


fda_path = SOURCE_DIR / SOURCES["FDA2007"][0]
with pdfplumber.open(fda_path) as pdf:
    appendix_pages = [i + 1 for i, p in enumerate(pdf.pages)
                      if "TARGET PRODUCT PROFILE TEMPLATE" in (p.extract_text() or "")
                      and "APPENDIX C" in (p.extract_text() or "")]
    n_pages = len(pdf.pages)
    first_appendix_page = max(appendix_pages)   # page 3 is the TOC entry
    milestone_tables = pdf.pages[first_appendix_page - 1].extract_tables()

lines = fda_lines(fda_path, first_appendix_page, n_pages)

sections = []
current = None
for line in lines:
    words = line["words"]
    first = words[0]
    if 66 <= first["x0"] <= 80 and FDA_SECTION_RE.match(first["text"]) and len(words) > 1:
        current = {"number": first["text"], "title_words": words[1:], "page": line["page"],
                   "body": [], "annot_x": None, "state": "pre"}
        sections.append(current)
        continue
    if current is None:
        continue
    texts = [w["text"] for w in words]
    if current["state"] == "pre" and "Target" in texts and "Annotations" in texts:
        current["annot_x"] = [w for w in words if w["text"] == "Annotations"][0]["x0"]
        current["state"] = "columns"
        continue
    if texts[0] == "Comments:":
        current["state"] = "comments"
        rest = words[1:]
        if rest:
            current["body"].append(("comments", text_of(rest)))
        continue
    if current["state"] == "columns":
        left = [w for w in words if w["x0"] < current["annot_x"] - 5]
        right = [w for w in words if w["x0"] >= current["annot_x"] - 5]
        if left:
            current["body"].append(("target", text_of(left)))
        if right:
            current["body"].append(("annotation", text_of(right)))
    elif current["state"] == "comments":
        current["body"].append(("comments", text_of(words)))
    else:
        current["body"].append(("guidance", text_of(words)))

fda_rows = []
for s in sections:
    title_text = text_of(s["title_words"])
    is_sub = "." in s["number"]
    if is_sub and ":" in title_text:
        title, first_guidance = title_text.split(":", 1)
        guidance_parts = [first_guidance.strip()]
    elif " — " in title_text:
        title, first_guidance = title_text.split(" — ", 1)
        guidance_parts = [first_guidance.strip()]
    else:
        title, guidance_parts = title_text, []
    guidance_parts += [t for kind, t in s["body"] if kind == "guidance"]
    fda_rows.append({
        "source": "FDA2007",
        "element_id": s["number"],
        "group": None if not is_sub else s["number"].split(".")[0],
        "element_label": title.strip(),
        "element_kind": "label_subsection" if is_sub else "label_section",
        "columns": "Target | Annotations | Comments" if s["annot_x"] is not None else "",
        "target_prompt": " ".join(t for kind, t in s["body"] if kind == "target"),
        "annotation_prompt": " ".join(t for kind, t in s["body"] if kind == "annotation"),
        "guidance_text": " ".join(p for p in guidance_parts if p),
        "page": s["page"],
    })
fda = pd.DataFrame(fda_rows)

milestone_rows = []
for table in milestone_tables:
    n_header = next((i for i, r in enumerate(table) if (r[0] or "").strip()), len(table))
    if n_header == 0:
        continue
    combined = [" ".join((r[c] or "").strip() for r in table[:n_header] if c < len(r) and (r[c] or "").strip())
                for c in range(max(len(r) for r in table[:n_header]))]
    header = [h for h in combined if h]
    if not header or not header[0].startswith("Milestone"):
        continue
    for row in table[n_header:]:
        label = (row[0] or "").strip()
        if label:
            milestone_rows.append({"source": "FDA2007", "element_id": f"M{len(milestone_rows) + 1}",
                                   "group": "Milestone box", "element_label": label,
                                   "element_kind": "milestone",
                                   "columns": " | ".join(h.replace("\\n", " ") for h in header[1:]),
                                   "page": first_appendix_page})
fda = pd.concat([pd.DataFrame(milestone_rows), fda], ignore_index=True)
print(f"Appendix C pages {first_appendix_page}-{n_pages}; milestones: {len(milestone_rows)}; "
      f"sections: {(fda.element_kind == 'label_section').sum()}; subsections: {(fda.element_kind == 'label_subsection').sum()}")
fda[["element_id", "element_kind", "element_label", "columns"]]
"""))

cells.append(markdown_cell("""
### FDA label sections → SPL section LOINC codes

The codes are copied from FDA's *Section Headings (LOINC)* page for Structured Product Labeling (https://www.fda.gov/industry/structured-product-labeling/section-headings-loinc, content current as of 06/23/2022). Only main sections are mapped. Subsection codes were not retrieved and are **not** filled. The 2007 section 16 *How Supplied/Storage and Handling* maps to two FDA headings, *How Supplied* (34069-5) and *Storage and Handling* (44425-7); both are kept.

Context, not extracted: the 2007 template's subsections 8.1–8.3 (pregnancy categories A/B/C/D/X, Labor and Delivery, Nursing Mothers) predate the 2014 Pregnancy and Lactation Labeling Rule, which restructured them. The 2007 template is therefore not current label structure in section 8.
"""))

cells.append(code_cell("""
SPL_LOINC = {   # from FDA "Section Headings (LOINC)" page, content current as of 06/23/2022
    "Indications and Usage": "34067-9",
    "Dosage and Administration": "34068-7",
    "Dosage Forms and Strengths": "43678-2",
    "Contraindications": "34070-3",
    "Warnings and Precautions": "43685-7",
    "Adverse Reactions": "34084-4",
    "Drug Interactions": "34073-7",
    "Use in Specific Populations": "43684-0",
    "Drug Abuse and Dependence": "42227-9",
    "Overdosage": "34088-5",
    "Description": "34089-3",
    "Clinical Pharmacology": "34090-1",
    "Nonclinical Toxicology": "43680-8",
    "Clinical Studies": "34092-7",
    "References": "34093-5",
    "How Supplied/Storage and Handling": "34069-5 | 44425-7",
    "Patient Counseling Information": "88436-1",
}

fda["spl_loinc"] = fda.apply(
    lambda r: SPL_LOINC.get(r["element_label"], "") if r["element_kind"] == "label_section" else "", axis=1)
unmapped = fda[(fda.element_kind == "label_section") & (fda.spl_loinc == "")]
if not unmapped.empty:
    raise ValueError(f"FDA sections without LOINC mapping: {unmapped.element_label.tolist()}")
fda[fda.element_kind == "label_section"][["element_id", "element_label", "spl_loinc"]]
"""))

cells.append(markdown_cell("""
## WOAH 2026 and NIH SEED — pdfplumber tables

Both sources are real PDF tables, but `extract_tables()` scatters cell text across grid columns. Rules:

- A row containing `Attributes` is the **column header**. Header names and their grid positions are kept.
- A row whose only non-empty cell sits at grid position ≤ 1 is a **group header**, e.g. `1. Product use summary` or `Preclinical Studies`. The exception is a numbered attribute such as `1.3. Further comments`, which is an attribute with an empty Comments cell. The first row of a table without a column header is also a group header; its cell fragments are joined, which repairs e.g. `Regulatory Guideline` + `s`.
- A row whose only non-empty cell sits at position ≥ 3 is a **continuation** of the previous row's value.
- A row containing `:` in its first cells (`Product Type: …`, `Indication: …`) is **document metadata**.
- Otherwise the first non-empty cell is the **attribute label**, and each further cell is assigned to the nearest header column. Tables without a column header have one value column, `Description`.
- The table title row (`WOAH TPP for a [disease name] vaccine`) is recorded as metadata.
"""))

cells.append(code_cell("""
def clean(cell):
    return re.sub(r"\\s+", " ", (cell or "").replace("\\n", " ")).strip()


def parse_plumber_tables(source, pdf_path, filled_example):
    rows, meta = [], {}
    header = None          # list of (position, name)
    group = None
    last = None
    with pdfplumber.open(pdf_path) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables():
                header_seen_in_table = False
                for r_idx, raw in enumerate(table):
                    cells = [(i, clean(c)) for i, c in enumerate(raw) if clean(c)]
                    if not cells:
                        continue
                    texts = [t for _, t in cells]
                    if "Attributes" in texts:
                        header = [(i, t) for i, t in cells if t != "Attributes"]
                        header_seen_in_table = True
                        continue
                    if any(":" in t for t in texts[:2]) and r_idx == 0:
                        for t in texts:
                            k, v = t.split(":", 1)
                            meta[k.strip()] = v.strip()
                        continue
                    if r_idx == 0 and not header_seen_in_table and (len(cells) > 1 or cells[0][0] <= 1) \\
                            and not re.match(r"^\\d+\\.\\d+\\.", texts[0]):
                        title = "".join(texts) if len(cells) > 1 else texts[0]
                        if title.startswith("WOAH TPP for"):
                            meta["table_title"] = title
                        else:
                            group = title
                            header = None
                        continue
                    if len(cells) == 1 and cells[0][0] <= 1 and not re.match(r"^\\d+\\.\\d+\\.", texts[0]):
                        group = texts[0]
                        continue
                    if len(cells) == 1 and cells[0][0] >= 3 and last is not None:
                        col = list(last["values"].keys())[-1]
                        last["values"][col] += " " + texts[0]
                        continue
                    label = texts[0]
                    values = {}
                    for pos, t in cells[1:]:
                        if header:
                            name = min(header, key=lambda h: abs(h[0] - pos))[1]
                        else:
                            name = "Description"
                        values[name] = (values.get(name, "") + " " + t).strip()
                    last = {"source": source, "group": group, "element_label": label,
                            "values": values, "page": pno,
                            "columns": " | ".join(h[1] for h in header) if header else "Description"}
                    rows.append(last)
    out = []
    for i, r in enumerate(rows, start=1):
        m = re.match(r"^(\\d+\\.\\d+)\\.\\s+(.*)$", r["element_label"])
        element_id = m.group(1) if m else str(i)
        label = m.group(2) if m else r["element_label"]
        row = {"source": source, "element_id": element_id, "group": r["group"], "element_label": label,
               "element_kind": "attribute", "columns": r["columns"], "page": r["page"]}
        for k, v in r["values"].items():
            key = ("example_value_" if filled_example and k != "Comments" else "guidance_") + \\
                re.sub(r"[^a-z0-9]+", "_", k.lower()).strip("_")
            row[key] = v
        out.append(row)
    return pd.DataFrame(out), meta


woah, woah_meta = parse_plumber_tables("WOAH2026", SOURCE_DIR / SOURCES["WOAH2026"][0], filled_example=False)
print("WOAH metadata:", woah_meta)
woah[["element_id", "group", "element_label", "columns", "guidance_comments"]]
"""))

cells.append(code_cell("""
seed_frames, seed_meta = [], {}
for key in ["SEED_SMALL_MOLECULE", "SEED_BIOLOGIC", "SEED_VACCINE", "SEED_CGT"]:
    df, meta = parse_plumber_tables(key, SOURCE_DIR / SOURCES[key][0], filled_example=True)
    seed_frames.append(df)
    seed_meta[key] = meta
seed = pd.concat(seed_frames, ignore_index=True)
print(pd.DataFrame(seed_meta).T)
print()
print(pd.crosstab(seed["group"].fillna("(attribute table)"), seed["source"]))
"""))

cells.append(markdown_cell("""
## NCATS PaVe-GT — Word table

A single table, *Product Target | Minimum Acceptable Result | Ideal Result*, read per `w:tc`. In the source, *Mechanism of Action* has two separate cells with identical text; they are not merged. `minimum_equals_ideal` records where the two cells hold the same text.
"""))

cells.append(code_cell("""
pave_doc = docx.Document(str(SOURCE_DIR / SOURCES["PAVEGT"][0]))
if len(pave_doc.tables) != 1:
    raise ValueError(f"Expected 1 table in PaVe-GT template, found {len(pave_doc.tables)}")
table = pave_doc.tables[0]
trs = table._tbl.tr_lst
header = [_Cell(tc, table).text.strip() for tc in trs[0].tc_lst]
if header != ["Product Target", "Minimum Acceptable Result", "Ideal Result"]:
    raise ValueError(f"Unexpected PaVe-GT header: {header}")
pave_rows = []
for i, tr in enumerate(trs[1:], start=1):
    tcs = tr.tc_lst
    texts = [clean(_Cell(tc, table).text) for tc in tcs]
    if len(texts) != 3:
        raise ValueError(f"PaVe-GT row {i} has {len(texts)} cells: {texts}")
    row = {"source": "PAVEGT", "element_id": str(i), "group": None, "element_label": texts[0],
           "element_kind": "attribute", "columns": " | ".join(header[1:]), "page": None,
           "guidance_minimum_acceptable_result": texts[1],
           "guidance_ideal_result": texts[2],
           "minimum_equals_ideal": texts[1] == texts[2]}
    pave_rows.append(row)
pave = pd.DataFrame(pave_rows)
pave_meta = {p.text.split(":")[0].strip(): p.text.split(":", 1)[1].strip()
             for p in pave_doc.paragraphs if p.text.startswith(("Product name:", "Therapeutic Modality:"))}
print("PaVe-GT metadata:", pave_meta)
pave[["element_id", "element_label", "minimum_equals_ideal", "guidance_minimum_acceptable_result"]]
"""))

cells.append(markdown_cell("""
## Combine into one long table and compare column models

The **column model** is what distinguishes TPP families:

| Role | FDA2007 | WOAH2026 | PAVEGT | SEED |
|---|---|---|---|---|
| target / threshold | `Target` (labeling language) | `Ideal`, `Minimum` | `Ideal Result`, `Minimum Acceptable Result` | `Preferred`, `Minimal Acceptance Criteria` |
| evidence link | `Annotations` (Protocol #, Serial #, Submission date) | — | — | — |
| commentary | `Comments` | `Comments` (delete when done) | — | — |
| versioning | milestone box (meeting × TPP version date) | — | — | — |

The table below is computed from the extracted `columns` values, not from this note.
"""))

cells.append(code_cell("""
elements = pd.concat([fda, woah, pave, seed], ignore_index=True, sort=False)
elements["norm_label"] = elements["element_label"].str.lower().str.replace(r"[^a-z0-9]+", " ", regex=True).str.strip()

column_models = (elements.groupby(["source", "columns"]).size()
                 .rename("n_elements").reset_index())
column_models
"""))

cells.append(markdown_cell("""
## Attribute inventory across sources

Rows are normalised labels (lower case, punctuation removed); cells give each source's element id. The only matching is **exact normalised label equality**. Related concepts with different wording (e.g. FDA *Dosage and Administration* vs SEED *Dosage Frequency* / *Route of Administration*) are deliberately **not** merged. That crosswalk is a modelling decision and is left for review in `crosswalk_review.csv`, which has an empty `concept` column to fill in.
"""))

cells.append(code_cell("""
inv = elements[elements["element_kind"].isin(["attribute", "label_section", "label_subsection"])].copy()
inv["source_family"] = inv["source"].str.replace(r"^SEED_.*$", "SEED", regex=True)
inventory = (inv.groupby(["norm_label", "source_family"])["element_id"]
             .apply(lambda s: ", ".join(dict.fromkeys(s.astype(str))))
             .unstack(fill_value=""))
inventory["n_families"] = (inventory != "").sum(axis=1)
inventory = inventory.sort_values(["n_families", "norm_label"], ascending=[False, True])
print(f"{len(inventory)} distinct labels; shared by >1 family: {(inventory.n_families > 1).sum()}")
inventory[inventory.n_families > 1]
"""))

cells.append(code_cell("""
crosswalk_review = (inv[["source_family", "source", "group", "element_id", "element_label", "columns"]]
                    .drop_duplicates(subset=["source_family", "group", "element_label"])
                    .assign(concept="", note="")
                    .sort_values(["source_family", "source", "element_id"]))
print(crosswalk_review.groupby("source_family").size().to_string())
"""))

cells.append(markdown_cell("""
## Write outputs

| File | Content |
|---|---|
| `tpp_elements.csv` | every extracted element (milestones, label sections/subsections, attributes) with columns, prompts/guidance, example values, LOINC |
| `tpp_column_models.csv` | column model per source |
| `tpp_label_inventory.csv` | normalised labels × source family |
| `crosswalk_review.csv` | one row per distinct element per family, with empty `concept` / `note` for manual crosswalk |
| `tpp_source_metadata.json` | document metadata found in the files (product type, indication, product name, table title) |
"""))

cells.append(code_cell("""
import json

elements.to_csv(OUTPUT_DIR / "tpp_elements.csv", index=False)
column_models.to_csv(OUTPUT_DIR / "tpp_column_models.csv", index=False)
inventory.to_csv(OUTPUT_DIR / "tpp_label_inventory.csv")
crosswalk_review.to_csv(OUTPUT_DIR / "crosswalk_review.csv", index=False)
with open(OUTPUT_DIR / "tpp_source_metadata.json", "w", encoding="utf-8") as f:
    json.dump({"sources": {k: {"file": v[0], "url": v[1]} for k, v in SOURCES.items()},
               "WOAH2026": woah_meta, "PAVEGT": pave_meta, "SEED": seed_meta},
              f, indent=1, ensure_ascii=False)
print(elements.groupby(["source", "element_kind"]).size().to_string())
for f in sorted(OUTPUT_DIR.iterdir()):
    print(f"{f.name:28s} {f.stat().st_size:>9,d} bytes")
"""))

for i, cell in enumerate(cells):
    cell["id"] = f"cell-{i:02d}"

notebook = {
    "cells": cells,
    "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
                 "language_info": {"name": "python"}},
    "nbformat": 4,
    "nbformat_minor": 5,
}
with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
    json.dump(notebook, f, indent=1, ensure_ascii=False)
print("wrote", OUTPUT_PATH, "cells:", len(cells))
