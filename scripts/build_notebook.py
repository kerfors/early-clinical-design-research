"""Build eupearl_template_extraction.ipynb with the json module.

Usage: python3 build_notebook.py <output.ipynb> <input_dir_for_notebook>
"""
import json
import sys


def markdown_cell(text):
    lines = text.strip("\n").split("\n")
    return {
        "cell_type": "markdown",
        "metadata": {},
        "source": [line + "\n" for line in lines],
    }


def code_cell(text):
    lines = text.strip("\n").split("\n")
    return {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [line + "\n" for line in lines],
    }


OUTPUT_PATH = sys.argv[1]
INPUT_DIR = sys.argv[2]

cells = []

cells.append(markdown_cell("""
# EU-PEARL platform trial templates — structural extraction

Extracts the structure of the five EU-PEARL templates (V4/V3/V1, 25 April 2023) into CSV and JSON:

| Role | Document |
|---|---|
| MPT | Master Protocol Template V4 (based on TransCelerate CPT v8.0) |
| ISA | Intervention Specific Appendix template V4 (based on CPT v8.0) |
| SAP | Statistical Analysis Plan template V3 (based on TransCelerate SAP template v3.0) |
| DMC | Data Monitoring Committee charter template V1 |
| CTR | Guidance for supplementary information to the CTR cover letter V4 |

Source: https://eu-pearl.eu/eu-pearls-suite-of-master-protocol-templates/

All values come from the .docx files. Nothing is inferred beyond the explicit rules documented below.
"""))

cells.append(markdown_cell("""
## Text conventions used for classification

Taken from the templates' own convention sections (MPT paragraphs 20 and 40–44; DMC charter paragraphs 96–98).

**CPT family (MPT, ISA, SAP)**

| Category | Rule in this notebook | Template convention |
|---|---|---|
| `instruction` | paragraph style `CPT_Instructional` | instructional text, aids authoring |
| `common` | between `<Start of common text>` and `<End of common text>` | harmonised language |
| `suggested` | between `<Start of suggested text>` / `<End …>` | adaptable language |
| `example` | between `<Start of example text>` / `<End …>`, or style `CPT_Example` | remove if not used |
| `sample` | between `<Start of sample text>` / `<End …>` | (one occurrence in MPT) |
| `marker` | paragraph that is only a start/end marker | — |
| variable spans | runs with character style `CPT_Variable` or colour `0070C0` | "blue bracketed text" |
| EU-PEARL change spans | runs with yellow highlight | "modifications … to the CPT V8 … for the master protocol approach" |

**DMC charter**

| Category | Rule | Convention |
|---|---|---|
| `guidance` | paragraph style starting `Guide` | guidance text (red) |
| optional spans | turquoise highlight | "bracketed / highlighted in blue … optional or modifiable" |
| fill-in spans | bright green highlight | "fill-in-the-blank or require an action" |

**CTR guidance**: no style conventions; paragraphs are `body`, the value is in the two overview tables.

Paragraphs not matching any rule are `body`.
"""))

cells.append(code_cell(f"""
import json
import re
from collections import Counter
from pathlib import Path

import docx
import pandas as pd
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

INPUT_DIR = Path("{INPUT_DIR}")
OUTPUT_DIR = Path("output_eupearl")
OUTPUT_DIR.mkdir(exist_ok=True)

pd.set_option("display.max_colwidth", 120)
pd.set_option("display.width", 200)
"""))

cells.append(markdown_cell("""
## Locate the five documents

Files are identified by name pattern, so the local file names (`1.EU-PEARL_MPT-Template_V4-25April2023.docx` etc.) and prefixed copies both work. Exactly one file per role is required.
"""))

cells.append(code_cell("""
ROLE_PATTERN = re.compile(r"EU-PEARL[_-](MPT|ISA|SAP|DMC|Guidance)", re.IGNORECASE)
ROLE_NAMES = {"MPT": "MPT", "ISA": "ISA", "SAP": "SAP", "DMC": "DMC", "GUIDANCE": "CTR"}
CPT_ROLES = {"MPT", "ISA", "SAP"}

DOC_PATHS = {}
for path in sorted(INPUT_DIR.glob("*.docx")):
    match = ROLE_PATTERN.search(path.name)
    if match is None:
        raise ValueError(f"Unrecognised file name: {path.name}")
    role = ROLE_NAMES[match.group(1).upper()]
    if role in DOC_PATHS:
        raise ValueError(f"Two files for role {role}: {DOC_PATHS[role].name}, {path.name}")
    DOC_PATHS[role] = path

missing = set(ROLE_NAMES.values()) - set(DOC_PATHS)
if missing:
    raise FileNotFoundError(f"Missing roles in {INPUT_DIR}: {sorted(missing)}")

for role, path in DOC_PATHS.items():
    print(f"{role}: {path.name}")
"""))

cells.append(markdown_cell("""
## Helper functions

- Body-order iteration over paragraphs and tables, so tables get the section they sit in.
- Runs are taken from `iter_inner_content()` so text inside hyperlinks is included.
- Heading styles are mapped explicitly. An unknown `Heading…` style raises an error instead of being guessed.
"""))

cells.append(code_cell("""
MARKER_RE = re.compile(r"<\\s*(Start|End)\\s+of\\s+(\\w+)\\s+text\\s*>", re.IGNORECASE)
LOOSE_MARKER_RE = re.compile(r"\\b(Start|End)\\s+of\\s+(\\w+)\\s+text\\s*>", re.IGNORECASE)
NUMBERED_HEADING_RE = re.compile(r"^Heading (\\d)$")
UNNUMBERED_HEADING_RE = re.compile(r"^Heading (\\d) \\((No TOC|non-numbered|No Numbering)\\)$")
FIELD_LABEL_STYLE = "Heading (No TOC)"
MANUAL_NUMBER_RE = re.compile(r"^\\d+(\\.\\d+)*\\.?\\s+")


def iter_body(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document._body)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document._body)


def iter_runs(paragraph):
    for item in paragraph.iter_inner_content():
        if hasattr(item, "runs"):
            yield from item.runs
        else:
            yield item


def run_rgb(run):
    color = run.font.color
    if color is None or color.type is None or color.rgb is None:
        return None
    return str(color.rgb)


def heading_info(style_name):
    \"\"\"Return (kind, level, numbered). kind: heading | field_label | None.\"\"\"
    if style_name == "Heading 0":
        return "heading", 1, False
    match = NUMBERED_HEADING_RE.match(style_name)
    if match:
        return "heading", int(match.group(1)), True
    match = UNNUMBERED_HEADING_RE.match(style_name)
    if match:
        return "heading", int(match.group(1)), False
    if style_name == FIELD_LABEL_STYLE:
        return "field_label", None, False
    if style_name.startswith("Heading"):
        raise ValueError(f"Unmapped heading style: {style_name}")
    return None, None, False


def numbering_suppressed(paragraph):
    ppr = paragraph._p.pPr
    if ppr is None or ppr.numPr is None or ppr.numPr.numId is None:
        return False
    return ppr.numPr.numId.val == 0


def joined_spans(runs, predicate):
    \"\"\"Concatenate contiguous runs matching predicate; return list of span texts.\"\"\"
    spans, current = [], []
    for run in runs:
        if predicate(run):
            current.append(run.text)
        elif current:
            spans.append("".join(current))
            current = []
    if current:
        spans.append("".join(current))
    return [s.strip() for s in spans if s.strip()]


def normalise_title(title):
    return re.sub(r"[^a-z0-9]", "", title.lower())


def cell_rows(table):
    \"\"\"Yield (row_idx, cell_idx, grid_span, v_merge, text) without merged-cell duplicates.\"\"\"
    for row_idx, tr in enumerate(table._tbl.tr_lst):
        for cell_idx, tc in enumerate(tr.tc_lst):
            tcpr = tc.tcPr
            grid_span = 1
            v_merge = None
            if tcpr is not None:
                span_el = tcpr.find(qn("w:gridSpan"))
                if span_el is not None:
                    grid_span = int(span_el.get(qn("w:val")))
                vm_el = tcpr.find(qn("w:vMerge"))
                if vm_el is not None:
                    v_merge = vm_el.get(qn("w:val")) or "continue"
            yield row_idx, cell_idx, grid_span, v_merge, _Cell(tc, table).text.strip()
"""))

cells.append(markdown_cell("""
## Extraction pass

One pass per document in body order:

- **Section tracking.** Numbered headings (`Heading 1`–`Heading 9`, unless numbering is switched off at paragraph level) get computed numbers such as `5.2.1`. Unnumbered headings (`… (No TOC)`, `(non-numbered)`, `Heading 0`) enter the path without a number. `Heading (No TOC)` paragraphs are **field labels** (title page and synopsis items such as `Sponsor Name:`) and do not open a section.
- **Text blocks.** Start/end markers are processed in sequence within each paragraph, using a stack. Mismatched or orphan end markers, blocks still open at the end (with their start paragraph), and malformed markers (e.g. a missing `<`) are recorded in `marker_events`, not repaired. Paragraphs after a broken marker therefore keep whatever category the stack gives them; review `marker_events` before relying on block categories near those positions.
"""))

cells.append(code_cell("""
def extract(role, path):
    document = docx.Document(str(path))
    counters = [0] * 9
    stack = []          # list of dicts: level, number, title
    block_stack = []
    para_rows, table_rows, section_rows, marker_events = [], [], [], []
    para_idx = -1
    table_idx = -1

    def current_section():
        if not stack:
            return {"section_number": None, "section_title": None, "section_path": None,
                    "h1_title": None, "section_key": None}
        top = stack[-1]
        return {
            "section_number": top["number"],
            "section_title": top["title"],
            "section_path": " > ".join((s["number"] + " " if s["number"] else "") + s["title"] for s in stack),
            "h1_title": stack[0]["title"],
            "section_key": top["key"],
        }

    for item in iter_body(document):
        if isinstance(item, Table):
            table_idx += 1
            sec = current_section()
            for row_idx, cell_idx, grid_span, v_merge, text in cell_rows(item):
                table_rows.append({"doc": role, "table_idx": table_idx, "after_para_idx": para_idx,
                                   **sec, "row_idx": row_idx, "cell_idx": cell_idx,
                                   "grid_span": grid_span, "v_merge": v_merge, "text": text})
            continue

        para_idx += 1
        paragraph = item
        text = paragraph.text
        style = paragraph.style.name
        if not text.strip():
            continue

        kind, level, numbered = heading_info(style)

        if kind == "heading":
            if numbered and numbering_suppressed(paragraph):
                numbered = False
            title = text.strip()
            number = None
            if numbered:
                counters[level - 1] += 1
                for i in range(level, 9):
                    counters[i] = 0
                number = ".".join(str(c) for c in counters[:level])
            while stack and stack[-1]["level"] >= level:
                stack.pop()
            key = f"{role}:{para_idx}"
            stack.append({"level": level, "number": number, "title": title, "key": key})
            section_rows.append({
                "doc": role, "section_key": key, "para_idx": para_idx, "level": level,
                "numbered": numbered, "number": number, "title": title,
                "title_clean": MANUAL_NUMBER_RE.sub("", title),
                "manual_number_in_title": bool(MANUAL_NUMBER_RE.match(title)),
                "style": style, "h1_title": stack[0]["title"],
                "parent_key": stack[-2]["key"] if len(stack) > 1 else None,
                "path": current_section()["section_path"],
            })
            category = "heading"
        elif kind == "field_label":
            category = "field_label"
        elif style.lower().startswith("toc"):
            category = "toc"
        else:
            # process block markers in sequence
            segments = MARKER_RE.split(text)
            # MARKER_RE has 2 groups -> split gives [text, kind, name, text, kind, name, ...]
            content_block = None
            content_seen = False
            only_markers = True

            def open_blocks():
                return "/".join(f"{n}@{s}" for n, s in block_stack)

            for i in range(0, len(segments), 3):
                seg_text = segments[i]
                if seg_text.strip():
                    only_markers = False
                    if not content_seen:
                        content_block = block_stack[-1][0] if block_stack else None
                        content_seen = True
                if i + 2 < len(segments):
                    action, name = segments[i + 1].lower(), segments[i + 2].lower()
                    open_names = [n for n, _ in block_stack]
                    if action == "start":
                        block_stack.append((name, para_idx))
                    elif block_stack and block_stack[-1][0] == name:
                        block_stack.pop()
                    elif name in open_names:
                        marker_events.append({"doc": role, "para_idx": para_idx, "event": "end_mismatch",
                                              "marker": name, "open_blocks": open_blocks(),
                                              "text": text.strip()[:120], **current_section()})
                        del block_stack[len(open_names) - 1 - open_names[::-1].index(name)]
                    else:
                        marker_events.append({"doc": role, "para_idx": para_idx, "event": "end_without_start",
                                              "marker": name, "open_blocks": open_blocks(),
                                              "text": text.strip()[:120], **current_section()})
            if MARKER_RE.search(text) and not only_markers:
                marker_events.append({"doc": role, "para_idx": para_idx, "event": "inline_marker",
                                      "marker": "", "open_blocks": open_blocks(),
                                      "text": text.strip()[:120], **current_section()})
            if len(LOOSE_MARKER_RE.findall(text)) > len(MARKER_RE.findall(text)):
                marker_events.append({"doc": role, "para_idx": para_idx, "event": "malformed_marker",
                                      "marker": "", "open_blocks": open_blocks(),
                                      "text": text.strip()[:120], **current_section()})

            if role in CPT_ROLES:
                if only_markers:
                    category = "marker"
                elif style == "CPT_Instructional":
                    category = "instruction"
                elif content_block is not None:
                    category = content_block
                elif style == "CPT_Example":
                    category = "example"
                else:
                    category = "body"
            elif role == "DMC":
                category = "guidance" if style.startswith("Guide") else "body"
            else:
                category = "body"

        runs = list(iter_runs(paragraph))
        row = {
            "doc": role, "para_idx": para_idx, **current_section(),
            "style": style, "category": category,
            "text": text.strip(),
            "hidden_chars": sum(len(r.text) for r in runs if r.font.hidden),
        }
        if role in CPT_ROLES:
            change_spans = joined_spans(runs, lambda r: r.font.highlight_color == WD_COLOR_INDEX.YELLOW)
            row["variable_spans"] = " | ".join(joined_spans(
                runs, lambda r: (r.style is not None and r.style.name == "CPT_Variable") or run_rgb(r) == "0070C0"))
            row["eupearl_change_spans"] = " | ".join(change_spans)
            row["eupearl_change_chars"] = sum(len(s) for s in change_spans)
        if role == "DMC":
            row["optional_spans"] = " | ".join(joined_spans(
                runs, lambda r: r.font.highlight_color == WD_COLOR_INDEX.TURQUOISE))
            row["fill_in_spans"] = " | ".join(joined_spans(
                runs, lambda r: r.font.highlight_color == WD_COLOR_INDEX.BRIGHT_GREEN))
        para_rows.append(row)

    for name, start_para in block_stack:
        marker_events.append({"doc": role, "para_idx": start_para, "event": "unclosed_at_end", "marker": name,
                              "open_blocks": "/".join(f"{n}@{s}" for n, s in block_stack),
                              "text": "", "section_number": None, "section_title": None,
                              "section_path": None, "h1_title": None, "section_key": None})
    return para_rows, table_rows, section_rows, marker_events
"""))

cells.append(code_cell("""
all_paras, all_tables, all_sections, all_markers = [], [], [], []
for role, path in DOC_PATHS.items():
    p, t, s, m = extract(role, path)
    all_paras += p
    all_tables += t
    all_sections += s
    all_markers += m

paragraphs = pd.DataFrame(all_paras)
tables = pd.DataFrame(all_tables)
sections = pd.DataFrame(all_sections)
markers = pd.DataFrame(all_markers)

print("non-empty paragraphs per doc:")
print(paragraphs.groupby("doc").size().to_string())
print()
print(pd.crosstab(paragraphs["doc"], paragraphs["category"]))
"""))

cells.append(markdown_cell("""
## Validate computed section numbers against each document's TOC

Each TOC entry (`toc N` style) holds the number Word rendered when the TOC was last updated, e.g. `1.1.<tab>Synopsis<tab>19`. Only headings within the TOC's depth are checked; deeper levels get `below_toc_depth`. A computed number that is not among the TOC numbers for the same title is reported, and so are headings with no TOC entry. The CTR guidance has no TOC.

**Known unresolved case — DMC charter.** The DMC TOC is stale: sections such as *Interim Analysis Group*, *Decision Making* and *New Intervention Meetings* are missing from it. The `ORGANIZATIONS` heading also has numbering switched off at paragraph level and a typed `.` instead. Three readings of the *Data Monitoring Committee* heading disagree: style-derived **3.1** (this notebook), TOC **4.1**, and LibreOffice rendering **2.1**, where hidden text around `BACKGROUND` appears to affect the result. DMC section numbers here are style-derived and **not confirmed**. Use titles, not numbers, when citing DMC sections.
"""))

cells.append(code_cell("""
TOC_LINE_RE = re.compile(r"^(\\d+(?:\\.\\d+)*)\\.?$")

toc_rows = []
for role, path in DOC_PATHS.items():
    document = docx.Document(str(path))
    for p in document.paragraphs:
        if not p.style.name.lower().startswith("toc") or not p.text.strip():
            continue
        parts = [x.strip() for x in p.text.split("\\t") if x.strip()]
        match = TOC_LINE_RE.match(parts[0])
        if match and len(parts) >= 2:
            number, title = match.group(1), parts[1]
        else:
            number, title = None, parts[0]
        toc_rows.append({"doc": role, "toc_number": number, "toc_title": title,
                         "norm_title": normalise_title(MANUAL_NUMBER_RE.sub("", title))})
toc = pd.DataFrame(toc_rows)

toc_depth = toc.dropna(subset=["toc_number"]).assign(
    depth=lambda df: df["toc_number"].str.count(r"\\.") + 1).groupby("doc")["depth"].max()

checks = []
for _, s in sections[sections["numbered"]].iterrows():
    candidates = toc[(toc["doc"] == s["doc"]) & (toc["norm_title"] == normalise_title(s["title_clean"]))]
    if s["doc"] not in toc_depth.index:
        status = "no_toc"
    elif s["level"] > toc_depth[s["doc"]]:
        status = "below_toc_depth"
    elif candidates.empty:
        status = "not_in_toc"
    elif s["number"] in set(candidates["toc_number"]):
        status = "ok"
    else:
        status = "number_differs"
    checks.append({"doc": s["doc"], "number": s["number"], "title": s["title"], "status": status,
                   "toc_numbers": ", ".join(str(x) for x in candidates["toc_number"])})
toc_check = pd.DataFrame(checks)
print(pd.crosstab(toc_check["doc"], toc_check["status"]))
toc_check[toc_check["status"] != "ok"]
"""))

cells.append(markdown_cell("""
## Section summary

One row per heading: content counts by category for the paragraphs **directly** under it (up to the next heading), the number of tables, and EU-PEARL change characters (yellow highlight) in the CPT-family documents.
"""))

cells.append(code_cell("""
direct = paragraphs[paragraphs["category"] != "heading"]
counts = pd.crosstab(direct["section_key"], direct["category"])
counts.columns = [f"n_{c}" for c in counts.columns]
section_summary = sections.merge(counts, left_on="section_key", right_index=True, how="left")
n_cols = [c for c in section_summary.columns if c.startswith("n_")]
section_summary[n_cols] = section_summary[n_cols].fillna(0).astype(int)

change = paragraphs.groupby("section_key")["eupearl_change_chars"].sum()
section_summary["eupearl_change_chars"] = section_summary["section_key"].map(change).fillna(0).astype(int)

n_tables = tables.groupby("section_key")["table_idx"].nunique()
section_summary["n_tables"] = section_summary["section_key"].map(n_tables).fillna(0).astype(int)

section_summary.groupby("doc").agg(sections=("section_key", "size"),
                                   numbered=("numbered", "sum"),
                                   change_chars=("eupearl_change_chars", "sum"))
"""))

cells.append(markdown_cell("""
## Field labels (title page and synopsis)

`Heading (No TOC)` paragraphs are labelled protocol data items, the CPT counterpart of M11 title-page and synopsis elements. A label's own paragraph may carry a bracketed variable, e.g. `Study Phase: [study phase]`.
"""))

cells.append(code_cell("""
labels = paragraphs[paragraphs["category"] == "field_label"].copy()
labels["label"] = labels["text"].str.split(":").str[0].str.replace(r"\\s+", " ", regex=True).str.strip()
labels["inline_value"] = labels["text"].str.split(":", n=1).str[1].str.strip()
field_labels = labels[["doc", "para_idx", "section_path", "label", "inline_value", "variable_spans"]]

label_matrix = (field_labels.assign(present=1)
                .pivot_table(index="label", columns="doc", values="present", aggfunc="sum", fill_value=0))
label_matrix
"""))

cells.append(markdown_cell("""
## MPT ↔ ISA section alignment

The ISA reuses the CPT backbone of the MPT. Sections are matched on the **exact normalised title path**, i.e. the section's title together with all its ancestor titles. Normalisation means lower case, non-alphanumerics removed, leading manual numbers stripped, and the `Appendix N:` prefix removed, because appendix numbers differ between MPT and ISA. There is no fuzzy matching; unmatched sections on either side are kept with `match = mpt_only / isa_only` for review.

For each ISA section, three flags are set **only where an explicit phrase occurs** in its direct paragraphs. The matched sentence is kept as evidence:

| Flag | Pattern (case-insensitive) |
|---|---|
| `refers_to_master` | `refer(s) to the master protocol`, `Section N.N of the master protocol` (section numbers kept in `master_section_refs`) |
| `extends_master` | `must be eligible according to the master protocol`, `add(itional) … intervention cohort specific`, `additional … criteria for intervention cohort` |
| `either_location` | `may be in the master protocol or the ISA`, `if (this content is / not) included in the master protocol` |

A section with none of the flags is labelled `no_explicit_master_reference`. That does **not** mean the content is ISA-specific.
"""))

cells.append(code_cell("""
MASTER_SECTION_REF_RE = re.compile(r"\\bsection\\s+(\\d+(?:\\.\\d+)*)\\s+of the master protocol\\b", re.IGNORECASE)

RELATION_PATTERNS = {
    "refers_to_master": [r"\\brefers? to the master protocol\\b",
                         r"\\bsection\\s+\\d+(?:\\.\\d+)*\\s+of the master protocol\\b"],
    "extends_master": [r"\\bmust be eligible according to the master protocol\\b",
                       r"\\badd(?:itional)?\\b[^.]*\\bintervention[ -]cohort[ -]specific\\b",
                       r"\\badditional\\b[^.]*\\bcriteria for intervention cohort\\b"],
    "either_location": [r"\\bmay be in the master protocol or the ISA\\b",
                        r"\\bif (?:this content is|not included) in the master protocol\\b"],
}
SENTENCE_SPLIT = re.compile(r"(?<=[.!?])\\s+")


def relation_flags(section_key):
    texts = paragraphs[(paragraphs["section_key"] == section_key)
                       & (~paragraphs["category"].isin(["heading", "field_label", "marker"]))]["text"]
    result = {flag: False for flag in RELATION_PATTERNS}
    evidence = []
    refs = []
    for text in texts:
        refs += MASTER_SECTION_REF_RE.findall(text)
        for sentence in SENTENCE_SPLIT.split(text):
            for flag, patterns in RELATION_PATTERNS.items():
                if any(re.search(p, sentence, re.IGNORECASE) for p in patterns):
                    result[flag] = True
                    evidence.append(f"[{flag}] {sentence.strip()}")
    result["evidence"] = " || ".join(evidence)
    result["master_section_refs"] = ", ".join(dict.fromkeys(refs))
    return result


APPENDIX_PREFIX_RE = re.compile(r"^appendix\\s+\\d+\\s*:\\s*", re.IGNORECASE)
title_by_key = section_summary.set_index("section_key")["title_clean"].to_dict()
parent_by_key = section_summary.set_index("section_key")["parent_key"].to_dict()


def align_title(title):
    return normalise_title(APPENDIX_PREFIX_RE.sub("", title))


def title_path_key(section_key):
    parts = []
    key = section_key
    while key is not None and not pd.isna(key):
        parts.append(align_title(title_by_key[key]))
        key = parent_by_key[key]
    return "|".join(reversed(parts))


def side(doc):
    df = section_summary[section_summary["doc"] == doc].copy()
    df["align_key"] = df["section_key"].map(title_path_key)
    return df


mpt_side = side("MPT")
isa_side = side("ISA")
dups = [d for d in (mpt_side, isa_side) if d["align_key"].duplicated().any()]
if dups:
    raise ValueError("Duplicate alignment keys: " + str([d.loc[d["align_key"].duplicated(keep=False), ["doc", "title"]].values.tolist() for d in dups]))

keep = ["align_key", "section_key", "number", "level", "title", "eupearl_change_chars"] + \\
    [c for c in section_summary.columns if c.startswith("n_")]
alignment = mpt_side[keep].merge(isa_side[keep], on="align_key", how="outer", suffixes=("_mpt", "_isa"))
alignment["match"] = alignment.apply(
    lambda r: "both" if pd.notna(r["section_key_mpt"]) and pd.notna(r["section_key_isa"])
    else ("mpt_only" if pd.notna(r["section_key_mpt"]) else "isa_only"), axis=1)

flags = alignment["section_key_isa"].map(lambda k: relation_flags(k) if pd.notna(k) else {})
flag_df = pd.DataFrame(flags.tolist(), index=alignment.index)
alignment = pd.concat([alignment, flag_df], axis=1)


def relation_label(r):
    if r["match"] == "mpt_only":
        return ""
    names = [f for f in RELATION_PATTERNS if r.get(f) is True]
    return "+".join(names) if names else "no_explicit_master_reference"


alignment["isa_relation"] = alignment.apply(relation_label, axis=1)
alignment = alignment.sort_values(["section_key_mpt", "section_key_isa"],
                                  key=lambda s: s.str.split(":").str[1].astype(float)).reset_index(drop=True)
print(alignment["match"].value_counts().to_string())
print()
print(alignment.loc[alignment["match"] != "mpt_only", "isa_relation"].value_counts().to_string())
"""))

cells.append(code_cell("""
alignment[alignment["match"] != "both"][["match", "number_mpt", "title_mpt", "number_isa", "title_isa"]]
"""))

cells.append(markdown_cell("""
## Where EU-PEARL changed CPT v8 for the master protocol approach

The MPT sections with the most yellow-highlighted characters, i.e. text EU-PEARL added or changed relative to TransCelerate CPT v8.0.
"""))

cells.append(code_cell("""
(section_summary[(section_summary["doc"] == "MPT") & (section_summary["eupearl_change_chars"] > 0)]
 .sort_values("eupearl_change_chars", ascending=False)
 [["number", "title", "eupearl_change_chars", "n_instruction"]].head(25))
"""))

cells.append(markdown_cell("""
## Marker anomalies

Inline markers, end markers without a matching start, and blocks still open at the end of the document. These are reported as found and not corrected.
"""))

cells.append(code_cell("""
markers[["doc", "para_idx", "event", "marker", "open_blocks", "section_number", "section_title", "text"]]
"""))

cells.append(markdown_cell("""
### Paragraphs affected by an unclosed block

For each block still open at the end of a document, list the later paragraphs outside the start section that carry that block's category. Their category comes from the broken block and is **not reliable**. It is left as extracted, for human review.
"""))

cells.append(code_cell("""
affected = []
for _, ev in markers[markers["event"] == "unclosed_at_end"].iterrows():
    start = paragraphs[(paragraphs["doc"] == ev["doc"]) & (paragraphs["para_idx"] == ev["para_idx"])].iloc[0]
    later = paragraphs[(paragraphs["doc"] == ev["doc"])
                       & (paragraphs["para_idx"] > ev["para_idx"])
                       & (paragraphs["category"] == ev["marker"])
                       & (paragraphs["section_key"] != start["section_key"])]
    affected.append(later.assign(unclosed_block_start=ev["para_idx"]))
unreliable_block_category = pd.concat(affected) if affected else paragraphs.iloc[0:0]
print(unreliable_block_category.groupby(["doc", "unclosed_block_start"]).size().to_string())
unreliable_block_category[["doc", "para_idx", "section_number", "section_title", "category", "text"]]
"""))

cells.append(markdown_cell("""
## CTR cover-letter overview tables

The CTR guidance gives two example overview tables, adapted from the CTFG 2019 recommendation paper: one trial with sub-protocols/ISAs added as substantial amendments, and separate trials sharing a master protocol. The tables are shown here as a grid, rebuilt from the cell rows. Horizontally merged cells appear once, with their `grid_span`.
"""))

cells.append(code_cell("""
ctr_tables = tables[tables["doc"] == "CTR"]
for t_idx, t in ctr_tables.groupby("table_idx"):
    grid = t.pivot_table(index="row_idx", columns="cell_idx", values="text", aggfunc="first").fillna("")
    print(f"--- CTR table {t_idx} ({grid.shape[0]} rows x {grid.shape[1]} cells)")
    print(grid.to_string(max_colwidth=40))
    print()
"""))

cells.append(markdown_cell("""
## Write outputs

| File | Content |
|---|---|
| `paragraphs.csv` | every non-empty body paragraph with section, category and spans |
| `sections.csv` | heading tree with direct-content counts |
| `toc_check.csv` | computed numbers against the documents' TOC |
| `field_labels.csv` | title page and synopsis labels |
| `mpt_isa_alignment.csv` | MPT↔ISA section matching with ISA relation flags and evidence |
| `tables.csv` | all table cells (long form) with section context |
| `marker_events.csv` | marker anomalies |
| `unreliable_block_category.csv` | paragraphs whose block category comes from an unclosed block |
| `eupearl_structure.json` | nested section tree per document |
"""))

cells.append(code_cell("""
paragraphs.to_csv(OUTPUT_DIR / "paragraphs.csv", index=False)
section_summary.to_csv(OUTPUT_DIR / "sections.csv", index=False)
toc_check.to_csv(OUTPUT_DIR / "toc_check.csv", index=False)
field_labels.to_csv(OUTPUT_DIR / "field_labels.csv", index=False)
alignment.to_csv(OUTPUT_DIR / "mpt_isa_alignment.csv", index=False)
tables.to_csv(OUTPUT_DIR / "tables.csv", index=False)
markers.to_csv(OUTPUT_DIR / "marker_events.csv", index=False)
unreliable_block_category.to_csv(OUTPUT_DIR / "unreliable_block_category.csv", index=False)

relation_by_isa_key = alignment.dropna(subset=["section_key_isa"]).set_index("section_key_isa")["isa_relation"]


def build_tree(doc):
    rows = section_summary[section_summary["doc"] == doc]
    nodes = {}
    roots = []
    for _, r in rows.iterrows():
        node = {
            "key": r["section_key"], "number": r["number"], "level": int(r["level"]),
            "numbered": bool(r["numbered"]), "title": r["title"], "style": r["style"],
            "para_idx": int(r["para_idx"]),
            "counts": {c[2:]: int(r[c]) for c in n_cols if r[c]},
            "eupearl_change_chars": int(r["eupearl_change_chars"]),
            "n_tables": int(r["n_tables"]),
            "children": [],
        }
        if doc == "ISA":
            node["isa_relation"] = relation_by_isa_key.get(r["section_key"], "")
        nodes[r["section_key"]] = node
        if r["parent_key"] is None or pd.isna(r["parent_key"]):
            roots.append(node)
        else:
            nodes[r["parent_key"]]["children"].append(node)
    return roots


structure = {
    "source": "https://eu-pearl.eu/eu-pearls-suite-of-master-protocol-templates/",
    "documents": {role: {"file": path.name, "sections": build_tree(role)} for role, path in DOC_PATHS.items()},
}
with open(OUTPUT_DIR / "eupearl_structure.json", "w", encoding="utf-8") as f:
    json.dump(structure, f, indent=1, ensure_ascii=False)

for f in sorted(OUTPUT_DIR.iterdir()):
    print(f"{f.name:28s} {f.stat().st_size:>10,d} bytes")
"""))

for i, cell in enumerate(cells):
    cell["id"] = f"cell-{i:02d}"

notebook = {
    "cells": cells,
    "metadata": {
        "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
        "language_info": {"name": "python"},
    },
    "nbformat": 4,
    "nbformat_minor": 5,
}

with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
    json.dump(notebook, f, indent=1, ensure_ascii=False)
print("wrote", OUTPUT_PATH, "cells:", len(cells))
